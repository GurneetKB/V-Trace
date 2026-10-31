"""Build the final presentation deck (deliverables/V-Trace_Final_Presentation.pptx).

The design follows the first-presentation deck: dark title/closing slides with the teal wave artwork and the
V-Trace logo, white content slides with a bold title, a one-line subtitle over a thin teal rule, the NUS-ISS logo
top-right, pastel cards with coloured numbers, teal chevrons and a navy key-message box (Aptos Display / Aptos).

All numbers come from results/veri_eval.json (VeRi-776) and results/demo_eval.json (demo cameras); screenshots from
docs/screenshots/. Speaker notes hold the script for the recorded 10-15 minute video.

    python docs/build/make_deck.py
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
RES, SHOTS, FIGS = ROOT / "results", ROOT / "docs" / "screenshots", ROOT / "docs" / "figures"
ART = ROOT / "docs" / "build" / "assets"
OUT = ROOT / "deliverables" / "V-Trace_Final_Presentation.pptx"
LOGO_DARK = ROOT / "assets" / "vtrace_logo_dark.png"
NUS = ART / "nus_iss_logo.png"

# palette of the first-presentation deck
INK, BODY_C, MUTED, RULE, BORDER = "161B22", "2B333C", "5C6670", "4FB3AE", "D6DEE6"
TEAL, TEAL_D, GREEN, GREEN_D, NAVY_BOX = "12A7A1", "0E8782", "5DC55B", "3E8E3A", "15324B"
MINT_T = "B8E9E6"
BLUE, PEACH, LAV, MINT, LGREEN, GREY, CREAM = "EAF1F8", "F4E9E0", "E9E7F5", "E9F5F3", "EFF5E8", "F3F5F7", "FDF4E8"
N_TEAL, N_ORANGE, N_PURPLE = "4FA5B5", "D08A55", "7E6DB8"
HEAD, BODY = "Aptos Display", "Aptos"

V = json.loads((RES / "veri_eval.json").read_text())      # scripts/eval_veri.py
D = json.loads((RES / "demo_eval.json").read_text())      # scripts/demo_eval.py
R = D["reid"]


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, runs, size=16, color=BODY_C, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=6, name=None):
    """runs: str, or list of paragraphs; a paragraph is a str, a list of (text, {opts}) runs, or
    {"bullet": True, "text": ...}."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        bullet = isinstance(para, dict) and para.get("bullet")
        if isinstance(para, dict):
            para = para["text"]
        for j, (txt, o) in enumerate([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = ("•  " if bullet and j == 0 else "") + txt
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = rgb(o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.035):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label(shape, lines, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER, font=BODY):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = font, Pt(size if i == 0 else size - 2), bold and i == 0
        r.font.color.rgb = rgb(color)


def chevron(slide, x, y, w=0.32, h=0.5, color=TEAL):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def picture(slide, path, x, y, w=None, h=None):
    if not Path(path).exists():
        ph = box(slide, x, y, w or 4, h or 3, GREY)
        label(ph, f"[missing image]\n{Path(path).name}", 12, MUTED, bold=False)
        return ph
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def fit_picture(slide, path, x, y, w, h):
    from PIL import Image
    if not Path(path).exists():
        return picture(slide, path, x, y, w, h)
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    return picture(slide, path, x + (w - iw * s) / 2, y + (h - ih * s) / 2, iw * s, ih * s)


NO = [0]


def content_slide(title, subtitle, notes):
    s = prs.slides.add_slide(BLANK)
    NO[0] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(s, 0.45, 0.3, 10.6, 0.7, title, 30, INK, bold=True, font=HEAD, name="Title")
    text(s, 0.47, 0.98, 11, 0.3, subtitle, 12, BODY_C, name="Subtitle")
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.45), Inches(1.33), Inches(13.0), Inches(1.33))
    line.line.color.rgb = rgb(RULE)
    line.line.width = Pt(0.75)
    picture(s, NUS, 11.4, 0.16, w=1.85)
    text(s, 12.55, 7.05, 0.45, 0.25, str(NO[0]), 9, BODY_C, align=PP_ALIGN.RIGHT, name="Slide number")
    s.notes_slide.notes_text_frame.text = notes
    return s


def dark_slide(art, notes):
    s = prs.slides.add_slide(BLANK)
    NO[0] += 1
    picture(s, art, 0, 0, w=13.333, h=7.5)
    picture(s, NUS, 11.4, 0.16, w=1.85)
    s.notes_slide.notes_text_frame.text = notes
    return s


def key_box(slide, x, y, w, h, msg, size=17):
    b = box(slide, x, y, w, h, NAVY_BOX, line=None, radius=0.12)
    label(b, msg, size, "FFFFFF", font=HEAD)


def style_chart(chart, colors, legend=True, fmt="0", size=12):
    chart.font.name, chart.font.size = BODY, Pt(size)
    chart.font.color.rgb = rgb(BODY_C)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
    for s, c in zip(chart.plots[0].series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.color.rgb = Pt(size - 1), rgb(BODY_C)
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E6EBEF")
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size - 1)
    chart.category_axis.format.line.color.rgb = rgb(BORDER)
    chart.category_axis.tick_labels.font.size = Pt(size)


def add_chart(slide, kind, cats, series, x, y, w, h, colors, **kw):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, [round(v, 1) for v in vals])
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    style_chart(chart, colors, **kw)
    return chart


# ================================================================== 1 title
s = dark_slide(ART / "bg_title.png",
    "Hello, I'm Gurneet Kaur Bhuller from Group 21. This is V-Trace, a vehicle search and re-identification platform. "
    "As in my proposal, the problem is finding one vehicle across several cameras when the plate is unreadable and the "
    "search starts from a description. Today I'll show what I built – V-Trace, using only pretrained, published models – "
    "how it performs on the VeRi-776 benchmark and on multi-camera street footage, what I learned, and then the live app.")
picture(s, LOGO_DARK, 0.55, 1.15, w=8.6)
text(s, 0.75, 4.25, 9, 0.5, "DESCRIBE  →  RANK  →  SELECT  →  TRACE", 22, GREEN, bold=True, font=HEAD)
text(s, 0.75, 5.05, 9, 0.5, "Multi-Attribute Vehicle Re-Identification System", 26, MINT_T)
text(s, 0.75, 5.95, 9, 0.4, "Gurneet Kaur Bhuller  |  Group 21", 17, "FFFFFF")
text(s, 0.75, 6.45, 9, 0.3, "Pattern Recognition Systems · Final Presentation", 12, "AEB8C2")

# ================================================================== 2 system design
s = content_slide("System Design",
                  "Every step uses a published, pretrained model – nothing is trained for this project.",
    "Here is the system. YOLO11 detects vehicles and ByteTrack follows each one; one track is one sighting, with four "
    "good crops and one full frame. Each sighting is described by three pretrained models: the FastReID re-identification "
    "model trained on VeRi-776, CLIP for colour, type and distinctive features, and EasyOCR for plates. The description "
    "search ranks sightings; after the user selects one, Re-ID similarity finds it in the other cameras.")
pipe = [("VIDEOS", "Demo or uploaded\n(one video = one camera)", BLUE),
        ("DETECT + TRACK", "YOLO11n + ByteTrack\n→ sightings", BLUE),
        ("DESCRIBE", "FastReID (VeRi-776)\nCLIP zero-shot\nEasyOCR", MINT),
        ("SEARCH", "User description\n→ ranked matches", "E6F2E9"),
        ("SELECT", "User picks the right vehicle", "FDF1E3"),
        ("Re-ID", "FastReID similarity\nthreshold + ratio test\n+ time window", "DDE7EC"),
        ("WHERE SEEN", "Location · real-world\ntime · frame per\nsighting · map", "E3F1EF")]
for i, (h, d, fill) in enumerate(pipe):
    x = 0.4 + i * 1.84
    box(s, x, 1.75, 1.6, 3.9, fill)
    text(s, x + 0.15, 1.93, 1.35, 0.8, h, 15, INK, bold=True, font=HEAD)
    text(s, x + 0.15, 3.0, 1.35, 2.5, d.split("\n"), 13, BODY_C, spacing=3)
    if i < 6:
        chevron(s, x + 1.63, 3.45, 0.18, 0.4)
text(s, 0.45, 6.05, 12.4, 0.6, "Sighting = one vehicle seen continuously in one camera · real-world time = camera recording start + position in video",
     14, INK, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.45, 6.45, 12.4, 0.4, "Streamlit app with admin sign-in: Search (describe → best matches → where seen) · Cameras (location + recording start) · How it works",
     12, MUTED, align=PP_ALIGN.CENTER)

# ================================================================== 3 pretrained models
s = content_slide("Pretrained Models Used",
                  "Each model is used exactly as published; V-Trace connects them into one workflow.",
    "These are the five pretrained components. YOLO11 nano detects cars, motorcycles, buses and trucks. ByteTrack links "
    "detections over time. CLIP compares each vehicle with text such as 'a photo of a red car' or 'a car with a rear "
    "spoiler', zero-shot. EasyOCR reads plates when they are large enough. And the key one: FastReID, the framework "
    "released by the VeRi-776 authors' team, with their model trained on VeRi-776. I load its published checkpoint and "
    "use it unchanged to turn every vehicle crop into a 2048-number signature.")
models = [("YOLO11n", "Vehicle detection\n(COCO: car · motorcycle · bus · truck)", "Ultralytics", BLUE, N_TEAL),
          ("ByteTrack", "Follows each vehicle while in view\n→ one sighting per vehicle", "Ultralytics", MINT, N_TEAL),
          ("CLIP ViT-B/32", "Zero-shot colour · type · features · free text", "OpenAI / open_clip", PEACH, N_ORANGE),
          ("EasyOCR", "Licence-plate text, wildcard matching (?)", "JaidedAI", CREAM, N_ORANGE),
          ("FastReID SBS(R50-ibn)", "Vehicle Re-ID, trained on VeRi-776\n2048-d signature per vehicle", "JDAI-CV · VeRi authors' team", LAV, N_PURPLE)]
for i, (h, d, src, fill, nc) in enumerate(models):
    x, y = (0.6 + (i % 3) * 4.15, 1.7) if i < 3 else (2.675 + (i - 3) * 4.15, 4.2)
    box(s, x, y, 3.85, 2.3, fill)
    text(s, x + 0.25, y + 0.2, 3.4, 0.45, h, 18, INK, bold=True, font=HEAD)
    text(s, x + 0.25, y + 0.75, 3.4, 1.0, d.split("\n"), 14, BODY_C, spacing=2)
    text(s, x + 0.25, y + 1.8, 3.4, 0.35, src, 12, nc, bold=True)
key_box(s, 0.6, 6.65, 12.15, 0.5, "NO TRAINING: PUBLISHED WEIGHTS, USED AS-IS", 15)

# ================================================================== 4 matching
s = content_slide("Search & Matching",
                  "A description score for step one; a cautious same-vehicle rule for step two.",
    "Two scoring steps. For the description, every field the user fills in gives each sighting a score between zero "
    "and one, and the final score is the weighted mean over only those fields, so leaving a field empty never hurts. "
    "For re-identification, the similarity of two sightings is the cosine of their FastReID signatures. A camera counts "
    "as 'seen' only if its best sighting is above 0.40, beats the camera's second-best by 0.08 – a ratio test, so "
    "look-alike cars are left to the user – and is within ten seconds in real-world time, which is each camera's "
    "recording start plus the position in its video.")
box(s, 0.6, 1.7, 5.6, 4.6, GREY)
text(s, 0.85, 1.9, 5.1, 0.4, "1  DESCRIPTION → RANKED MATCHES", 15, INK, bold=True, font=HEAD)
for i, (f, w) in enumerate((("plate (wildcards)", "0.35"), ("colour (CLIP)", "0.30"), ("type (CLIP)", "0.25"),
                            ("distinctive features", "0.10"), ("free text", "0.10"))):
    y = 2.45 + i * 0.55
    c = box(s, 0.85, y, 3.6, 0.44, "FFFFFF", radius=0.3)
    label(c, f, 13, BODY_C, bold=False)
    text(s, 4.6, y + 0.07, 1.2, 0.35, f"weight {w}", 12, TEAL_D, bold=True)
text(s, 0.85, 5.3, 5.1, 0.9, "Only the fields the user gives are scored – a missing field never counts against a vehicle.", 13)
box(s, 6.6, 1.7, 6.15, 4.6, "E3F1EF")
text(s, 6.85, 1.9, 5.6, 0.4, "2  SELECTED → OTHER CAMERAS", 15, INK, bold=True, font=HEAD)
text(s, 6.85, 2.45, 5.7, 3.8, [
    {"bullet": True, "text": [("Similarity ", {"bold": True}), ("= cosine of the FastReID signatures", {})]},
    {"bullet": True, "text": [("Threshold ≥ 0.40 ", {"bold": True}), ("– chosen from VeRi-776 score distributions", {})]},
    {"bullet": True, "text": [("Ratio test: ", {"bold": True}), ("best must beat the camera's 2nd best by 0.08 – look-alikes are left to the user", {})]},
    {"bullet": True, "text": [("Real-world time: ", {"bold": True}), ("camera recording start + position in the video; limit ≤ 10 s (5 s – 12 h)", {})]},
    {"bullet": True, "text": "No confident match → show the 3 closest vehicles instead"}], 15, spacing=8)
key_box(s, 0.6, 6.5, 12.15, 0.55, "EVERY RESULT COMES WITH ITS EVIDENCE: LOCATION · REAL-WORLD TIME · CROP · FRAME · SCORE", 15)

# ================================================================== 5 VeRi results
th = V[f"threshold_{R['threshold']}"]
same, diff = V["same_vehicle_similarity_percentiles_10_50_90"], V["different_vehicle_similarity_percentiles_50_99_99.9"]
s = content_slide("Results: Re-ID Model on VeRi-776",
                  "The published model reproduces its reported accuracy – and tells us where to set the threshold.",
    f"First, does the model work as published? On the official VeRi-776 split it reaches {V['rank1']:.1f} percent "
    f"rank-one and {V['mAP']:.1f} percent mean average precision; the FastReID model zoo reports 97.0 and 81.9. So the "
    f"weights load and behave correctly. The score distributions set the threshold: same-vehicle pairs from different "
    f"cameras have a median similarity of {same[1]:.2f}, while 99 percent of different-vehicle pairs score below "
    f"{diff[1]:.2f}. At 0.40, {th['same_vehicle_kept_%']:.0f} percent of true pairs pass and only "
    f"{th['different_vehicle_kept_%']:.1f} percent of wrong ones.")
add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ["Rank-1", "mAP"],
          [("Our run", [V["rank1"], V["mAP"]]), ("FastReID model zoo", [97.0, 81.9])], 0.6, 1.6, 7.4, 4.8, [TEAL, "B9C3CC"],
          fmt="0.0")
callouts = [(f"{V['rank1']:.1f}%", "Rank-1 (official split, no training)", BLUE, N_TEAL),
            (f"{V['mAP']:.1f}%", "mAP", MINT, TEAL_D),
            (f"{same[1]:.2f} vs {diff[1]:.2f}", "median same-vehicle vs 99th pct different-vehicle similarity", PEACH, N_ORANGE)]
for i, (v, l, fill, c) in enumerate(callouts):
    y = 1.7 + i * 1.55
    box(s, 8.4, y, 4.35, 1.38, fill)
    text(s, 8.65, y + 0.15, 3.9, 0.6, v, 28, c, bold=True, font=HEAD)
    text(s, 8.65, y + 0.85, 3.9, 0.45, l, 12, BODY_C)
key_box(s, 0.6, 6.5, 12.15, 0.55, f"THRESHOLD 0.40: {th['same_vehicle_kept_%']:.0f}% OF TRUE PAIRS PASS, "
                                  f"{th['different_vehicle_kept_%']:.1f}% OF WRONG PAIRS", 15)

# ================================================================== 6 demo results
w_, nw = R["real_time_window"], R["no_time_window"]
s = content_slide("Results: Multi-Camera Demo Footage",
                  "Real street footage is harder than the benchmark – a cautious rule keeps reported cameras mostly correct.",
    f"The demo footage is an openly licensed street recording cut into three camera zones along one road, so the same "
    f"cars really pass all three, plus a fourth camera on another road; tracking the full frame once gives a pseudo ground "
    f"truth. On this real street video the same model is much weaker: vehicles are small and seen from very different angles. The "
    f"correct sighting is ranked first for {R['Rank-1']:.0f} percent of queries. That is why the decision rule matters. "
    f"With the ratio test and a ten-second real-world time window, {w_['camera_precision_%']:.0f} percent of the cameras "
    f"V-Trace reports are correct. Part of that comes from the curated clocks: the other road was recorded minutes later, "
    f"so the window rules it out. Without the window, the model and ratio test alone reach {nw['camera_precision_%']:.0f} "
    f"percent, with {nw['queries_wrongly_seen_in_camD_%']:.0f} percent of queries wrongly placed on the other road. It "
    f"finds {w_['camera_recall_%']:.0f} percent of the cameras – for the rest it shows the closest vehicles so the user can decide.")
stats = [(f"{R['Rank-1']:.0f}% / {R['Rank-5']:.0f}%", "Rank-1 / Rank-5, other cameras", BLUE, N_TEAL),
         (f"{w_['camera_precision_%']:.0f}%", f"reported cameras correct, 10 s real-time window (none: {nw['camera_precision_%']:.0f}%)", MINT, TEAL_D),
         (f"{w_['camera_recall_%']:.0f}%", "of true cameras reported", LGREEN, GREEN_D),
         (f"{w_['queries_wrongly_seen_in_camD_%']:.0f}% / {nw['queries_wrongly_seen_in_camD_%']:.0f}%",
          "placed on the other road, with / without window", PEACH, N_ORANGE)]
for i, (v, l, fill, c) in enumerate(stats):
    y = 1.7 + i * 1.22
    box(s, 0.6, y, 3.6, 1.05, fill)
    text(s, 0.82, y + 0.12, 3.2, 0.5, v, 26, c, bold=True, font=HEAD)
    text(s, 0.82, y + 0.65, 3.2, 0.35, l, 12, BODY_C)
fit_picture(s, SHOTS / "seen.png", 4.55, 1.7, 8.2, 5.15)

# ================================================================== 7 demo
s = content_slide("Live Demo",
                  "Describe → best matches → select → where this vehicle was seen.",
    "Now the app. After the admin signs in, the four demo cameras are already processed. I describe the target: colour "
    "red, type SUV, and search. The ranked matches show red cars; number three is a red Citroën in camera B at 42 "
    "seconds past 8:15. I select it, and V-Trace shows where it was seen: Orchard Road west of Scotts Road at "
    "8:15:40.8, at the Scotts Road junction at 8:15:42.0 and east of it at 8:15:43.5, each with location, date and time, a crop and the full frame, plus a "
    "map. The other road, recorded minutes later, is listed as not seen. On the Cameras page each camera's location "
    "and recording start can be edited, and the route updates.")
for i, (f, cap) in enumerate((("search.png", "“red + suv” → ranked best matches → Select"),
                              ("seen.png", "Route with location, real-world time and a frame per camera"))):
    x = 0.6 + i * 6.2
    box(s, x, 1.7, 5.95, 4.45, GREY)
    fit_picture(s, SHOTS / f, x + 0.15, 1.85, 5.65, 4.15)
    text(s, x, 6.3, 5.95, 0.35, cap, 14, INK, bold=True, font=HEAD, align=PP_ALIGN.CENTER)

# ================================================================== 8 lessons
s = content_slide("Lessons Learned",
                  "Use what exists, verify it, and measure it on your own data.",
    f"Three lessons. First, a published model trained by the dataset's own authors beat anything I could train in ten "
    f"days, but I verified it by re-running the official benchmark before trusting it. Second, benchmark accuracy does "
    f"not transfer: on my street footage rank-one dropped to about {R['Rank-1']:.0f} percent. Third, decision rules matter "
    f"as much as scores: a ratio test and a time window raised the share of correct cameras from "
    f"{R['threshold_only']['camera_precision_%']:.0f} to {w_['camera_precision_%']:.0f} percent without any training.")
pit = [("USE WHAT EXISTS – AND CHECK IT", "FastReID's VeRi-776 model is far stronger than a model trained in ten days.",
        "Re-run the official benchmark first", f"Rank-1 {V['rank1']:.1f}% · mAP {V['mAP']:.1f}%", BLUE, N_TEAL),
       ("BENCHMARKS DON'T TRANSFER", "Small, distant vehicles and new angles in street footage.",
        "Measure on your own data (pseudo ground truth)", f"Rank-1 {V['rank1']:.0f}% → {R['Rank-1']:.0f}%", PEACH, N_ORANGE),
       ("DECISION RULES MATTER", "A threshold alone reports many look-alikes.",
        "Ratio test + time window; show candidates when unsure", f"Precision {R['threshold_only']['camera_precision_%']:.0f}% → {w_['camera_precision_%']:.0f}%", LAV, N_PURPLE)]
for i, (h, prob, fix, gain, fill, nc) in enumerate(pit):
    x = 0.6 + i * 4.15
    box(s, x, 1.7, 3.85, 4.0, fill)
    text(s, x + 0.25, 1.9, 3.4, 0.45, f"{i + 1:02d}", 20, nc, bold=True, font=HEAD)
    text(s, x + 0.25, 2.45, 3.4, 0.8, h, 17, INK, bold=True, font=HEAD)
    text(s, x + 0.25, 3.35, 3.4, 2.3, [[("Why  ", {"bold": True, "color": nc}), (prob, {})], "",
                                       [("So  ", {"bold": True, "color": TEAL_D}), (fix, {})]], 14, spacing=4)
    key_box(s, x, 5.9, 3.85, 0.8, gain, 16)

# ================================================================== 9 limitations
s = content_slide("Limitations & Next Steps",
                  "A focused proof of concept – the clearest next steps follow from the results.",
    "Limitations. The model was trained on VeRi images and is weaker on small, distant vehicles. Plates are rarely "
    "readable at street distance. Colour and type from zero-shot CLIP confuse similar classes. The demo cameras are zones "
    "of one recording, not separate physical cameras. And V-Trace reads recorded files, though it works frame by frame, "
    "so live streams are a small change.")
text(s, 0.75, 1.65, 5.5, 0.5, "LIMITATIONS", 20, "D45F5F", bold=True, font=HEAD)
text(s, 7.05, 1.65, 5.5, 0.5, "NEXT", 20, GREEN, bold=True, font=HEAD)
lim = [("Re-ID weaker on small, distant vehicles", "higher-resolution cameras · larger crops"),
       ("Plates unreadable at street distance", "dedicated plate detector before OCR"),
       ("Zero-shot colour / type confusions", "richer prompts · per-colour checks"),
       ("Demo cameras = zones of one recording", "separate physical cameras (e.g. CityFlow) · travel-time limits"),
       ("Recorded videos only", "live RTSP camera streams")]
for i, (l, n) in enumerate(lim):
    y = 2.3 + i * 0.82
    text(s, 0.75, y, 5.6, 0.6, "×  " + l, 17, BODY_C, anchor=MSO_ANCHOR.MIDDLE)
    b = box(s, 7.05, y, 5.7, 0.6, "EFF7F6")
    label(b, n, 14, INK, bold=False, align=PP_ALIGN.LEFT)
key_box(s, 2.3, 6.55, 8.75, 0.5, "MVP: recorded-video proof of concept  →  future extension: live camera streams", 14)

# ================================================================== 10 conclusion + thank you
s = dark_slide(ART / "bg_end.png",
    f"To conclude: V-Trace turns camera videos into searchable sightings, ranks them against a description, and once the "
    f"user selects the right vehicle, shows every camera it was seen in with times and frames. It is built entirely from "
    f"published pretrained models; the Re-ID model reproduces its benchmark accuracy, and on real street footage the "
    f"cameras it reports are correct {w_['camera_precision_%']:.0f} percent of the time. Thank you – I'm happy to take questions.")
picture(s, LOGO_DARK, 0.6, 0.45, w=4.2)
concl = [(f"{V['mAP']:.1f}%", "Re-ID mAP on VeRi-776, pretrained model as published"),
         ("0 models trained", "YOLO11 · ByteTrack · CLIP · EasyOCR · FastReID"),
         (f"{w_['camera_precision_%']:.0f}%", "of reported cameras correct on real street footage")]
for i, (v, l) in enumerate(concl):
    x = 0.75 + i * 3.9
    text(s, x, 2.0, 3.6, 0.7, v, 34, GREEN if i != 1 else "E39B57", bold=True, font=HEAD)
    text(s, x, 2.8, 3.4, 0.8, l, 14, "D9E2E8")
text(s, 0, 3.75, 13.333, 1.2, "Thank you!", 72, TEAL, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
text(s, 0.55, 6.25, 12.25, 1.0,
     "In the preparation of this project, AI-based tools including Claude, ChatGPT and Gemini were used in a supporting "
     "capacity to assist with tasks such as drafting and refining text, organising content, implementation and general "
     "research support. All ideas, analysis, technical design decisions and conclusions are the work of the project team.",
     11, "FFFFFF", bold=True)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"saved {OUT} ({NO[0]} slides)")
