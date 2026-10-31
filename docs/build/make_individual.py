"""Build the individual project report in the PRS template format (.docx + .pdf).

Template: "Master of Technology in AI Systems" line, centred title "Individual Project Report", a two-row table
(Your Name / Certificate) and the three numbered questions as blue headings. Text from docs/INDIVIDUAL_REPORT.md.

    python docs/build/make_individual.py
"""
import html
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "INDIVIDUAL_REPORT.md"
OUT = ROOT / "deliverables" / "V-Trace_Individual_Report_Gurneet_Kaur_Bhuller"
NAME = "Gurneet Kaur Bhuller (A0354164E)"
CERT = "Graduate Certificate in Pattern Recognition Systems"
HEAD_BLUE = RGBColor(0x36, 0x5F, 0x91)
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
AI_NOTE = ("AI declaration. In preparing this report and building the project, AI-based tools including Claude, "
           "ChatGPT and Gemini were used in a supporting capacity: drafting and refining text, organising content, "
           "and coding assistance whose output was reviewed and tested. All ideas, analysis, technical design "
           "decisions and conclusions are my own.")


def sections():
    """[(heading, [paragraph or ('bullet', text)])] from the markdown file."""
    out = []
    for block in re.split(r"^## ", SRC.read_text(encoding="utf-8"), flags=re.M)[1:]:
        head, *body = block.strip().split("\n")
        items, para = [], []
        for line in body + [""]:
            if line.startswith("- "):
                if para:
                    items.append(" ".join(para)); para = []
                items.append(("bullet", line[2:].strip()))
            elif line.startswith("  ") and items and isinstance(items[-1], tuple):
                items[-1] = ("bullet", items[-1][1] + " " + line.strip())
            elif line.strip():
                para.append(line.strip())
            elif para:
                items.append(" ".join(para)); para = []
        out.append((head.strip(), items))
    return out


def runs(par, text, size=11):
    """Add text with **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        r = par.add_run(part)
        r.bold = i % 2 == 1
        r.font.size = Pt(size)


def borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def docx_report(secs):
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(11)

    p = d.add_paragraph()
    r = p.add_run("Master of Technology in AI Systems"); r.font.size = Pt(14)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(12)
    r = p.add_run("Individual Project Report"); r.font.size = Pt(26)

    t = d.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for row, (k, v) in zip(t.rows, (("Your Name:", NAME), ("Certificate:", CERT))):
        row.height = Cm(1.1)
        row.cells[0].width, row.cells[1].width = Cm(5.6), Cm(11.0)
        for cell, txt in zip(row.cells, (k, v)):
            cell.paragraphs[0].add_run(txt).font.size = Pt(12)

    for head, items in secs:
        h = d.add_paragraph()
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(16), Pt(6)
        r = h.add_run(head)
        r.font.name = "Cambria"; r.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
        r.font.size, r.font.color.rgb = Pt(18), HEAD_BLUE
        for it in items:
            if isinstance(it, tuple):
                par = d.add_paragraph(style="List Bullet")
                runs(par, it[1])
            else:
                par = d.add_paragraph()
                runs(par, it)
            par.paragraph_format.space_after = Pt(4)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run(AI_NOTE); r.italic = True; r.font.size = Pt(9)
    d.save(OUT.with_suffix(".docx"))
    print("saved", OUT.with_suffix(".docx"))


def pdf_report(secs):
    def inline(t):
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(t))
    body = []
    for head, items in secs:
        body.append(f"<h2>{html.escape(head)}</h2>")
        bullets = []
        for it in items + [""]:
            if isinstance(it, tuple):
                bullets.append(f"<li>{inline(it[1])}</li>")
                continue
            if bullets:
                body.append("<ul>" + "".join(bullets) + "</ul>"); bullets = []
            if it:
                body.append(f"<p>{inline(it)}</p>")
    css = """
    @page { size: A4; margin: 18mm 20mm; }
    body { font-family: 'Times New Roman', serif; font-size: 11pt; line-height: 1.32; color: #111; }
    .prog { font-size: 14pt; margin: 0; }
    h1 { font-weight: normal; font-size: 26pt; text-align: center; margin: 10px 0 18px; }
    table { border-collapse: collapse; width: 92%; margin: 0 auto 6px; font-size: 12pt; }
    td { border: 1px solid #000; padding: 6px 10px 16px; vertical-align: top; }
    td:first-child { width: 32%; }
    h2 { font-family: Cambria, Georgia, serif; font-weight: normal; color: #365F91; font-size: 17pt; margin: 18px 0 6px; }
    p, li { text-align: justify; margin: 0 0 5px; }
    ul { margin: 0 0 6px; padding-left: 20px; }
    .ai { font-style: italic; font-size: 9pt; margin-top: 16px; border-top: 1px solid #999; padding-top: 6px; }
    """
    page = (f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>"
            f"<p class='prog'>Master of Technology in AI Systems</p><h1>Individual Project Report</h1>"
            f"<table><tr><td>Your Name:</td><td>{html.escape(NAME)}</td></tr>"
            f"<tr><td>Certificate:</td><td>{CERT}</td></tr></table>{''.join(body)}"
            f"<p class='ai'>{html.escape(AI_NOTE)}</p></body></html>")
    with tempfile.TemporaryDirectory() as tmp:
        f = Path(tmp) / "r.html"
        f.write_text(page, encoding="utf-8")
        subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--user-data-dir={Path(tmp) / 'p'}", f"--print-to-pdf={OUT.with_suffix('.pdf')}", f.as_uri()],
                       check=True, capture_output=True, timeout=180)
    print("saved", OUT.with_suffix(".pdf"))


if __name__ == "__main__":
    s = sections()
    docx_report(s)
    pdf_report(s)
